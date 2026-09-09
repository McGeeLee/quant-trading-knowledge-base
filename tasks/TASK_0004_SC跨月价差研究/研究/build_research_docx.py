from pathlib import Path
import json
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
v2=json.loads((ROOT/'output/v2_双腿回测明细_20260909.json').read_text())
v3=json.loads((ROOT/'output/v3_近月多_vs_远月多_明细.json').read_text())
d=Document()
s=d.sections[0]
s.page_width=Inches(8.5);s.page_height=Inches(11)
s.top_margin=s.bottom_margin=Inches(.7)
s.left_margin=s.right_margin=Inches(.75)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
 st=d.styles[name];st.font.name='Arial Unicode MS';st.font.color.rgb=RGBColor(0,0,0)
 fonts=st.element.get_or_add_rPr().get_or_add_rFonts()
 for attr in list(fonts.attrib):
  if 'Theme' in attr:del fonts.attrib[attr]
 fonts.set(qn('w:eastAsia'),'Arial Unicode MS')
 st.paragraph_format.space_after=Pt(4)
d.styles['Normal'].font.size=Pt(11)
d.styles['Normal'].paragraph_format.line_spacing=1.0
d.styles['Title'].font.size=Pt(23)
d.styles['Heading 1'].font.size=Pt(16)
d.styles['Heading 2'].font.size=Pt(12)
for style in d.styles:
 for border in style.element.findall('.//'+qn('w:pBdr')):
  border.getparent().remove(border)
d.styles['Subtitle'].font.italic=False
for name in ['Heading 1','Heading 2']:
 d.styles[name].paragraph_format.space_before=Pt(9)
d.core_properties.title='原油跨月价差与单腿做多策略研究'
d.core_properties.subject='SC 10月与11月合约 v2及v3历史回测比较'
d.core_properties.author=''

def p(t,style=None):return d.add_paragraph(t,style)
def h(t):d.add_heading(t,1)
def sub(t):d.add_heading(t,2)
def page():pass
def money(x):return f'{x*1000:,.0f}'
def table(headers,rows,widths):
 t=d.add_table(rows=1, cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 for c,w in zip(t.columns,widths):c.width=Inches(w)
 for c,txt in zip(t.rows[0].cells,headers):c.text=txt
 for row in rows:
  for c,txt in zip(t.add_row().cells,row):c.text=str(txt)
 for i,row in enumerate(t.rows):
  trpr=row._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
  if i==0:trpr.append(OxmlElement('w:tblHeader'))
  for j,c in enumerate(row.cells):
   c.width=Inches(widths[j]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   pr=c._tc.get_or_add_tcPr()
   sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'DCE6EE' if i==0 else 'FFFFFF');pr.append(sh)
   borders=OxmlElement('w:tcBorders')
   for edge in ['top','left','bottom','right']:
    e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
   pr.append(borders)
   margins=OxmlElement('w:tcMar')
   for side in ['top','bottom','left','right']:
    e=OxmlElement('w:'+side);e.set(qn('w:w'),'85');e.set(qn('w:type'),'dxa');margins.append(e)
   pr.append(margins)
   for para in c.paragraphs:
    para.alignment=WD_ALIGN_PARAGRAPH.LEFT if j==0 else WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.space_after=Pt(0);para.paragraph_format.line_spacing=1.1
    for run in para.runs:run.font.size=Pt(10);run.bold=i==0
 p('').paragraph_format.space_after=Pt(0)
 return t

p('原油跨月价差与单腿做多策略研究','Title')
p('SC 10月与11月合约  v2与v3历史回测比较','Subtitle')
p('整理日期  2026年9月9日    最新样本交易日  2026年9月8日')
h('研究结论')
p('现有样本能找到盈利交易，但尚不足以证明稳定、可实际成交的盈利优势。价差连续上涨三日后，v3近月单独做多的历史表现优于远月单独做多；这一优势依赖信号参数与流动性筛选，且单腿风险明显高于双腿价差交易。')
table(['策略','原始规则合计净收益','过滤版合计净收益'],[
 ['v2 多近月空远月','约0元','10,600元'],['v3 只做多近月','19,600元','65,300元'],['v3 只做多远月','−34,900元','−20,700元']],[2.4,2.3,2.3])
p('表中均为每腿1手、扣除假设成本并计入各样本期末未平仓盯市后的金额。v2往返成本为每组400元，v3为每笔200元；这是研究假设，不是逐日核实的实际费用。不同策略的风险、持有期和交易次数不同，不能仅按总利润选策略。')
sub('研究对象与问题')
p('SC2610和SC2611分别是2026年10月、11月交割月份的原油期货合约。研究将近月价格减远月价格定义为价差，检验“价差连续上涨三日”能否成为入场信号，并比较双腿对冲与单腿做多的结果。SC每手1000桶，价格变动1元每桶对应每手1000元盈亏。[1]')
sub('数据范围')
p('使用上期能源公开日行情配对数据，共1224个交易日，覆盖2020至2026年各年的10月和11月合约。2020年仅从6月17日起；其他年份大体从前一年12月初起，历史年份止于当年9月15日附近，2026年止于9月8日。[2]')
p('下文年份指合约年份，不是自然年度收益。例如2026年合约样本包含2025年12月行情。各年是独立研究窗口，未覆盖完整持有至到期过程，也不构成一个连续投资账户。')

page();h('策略规则与收益计算')
table(['项目','v2 双腿价差','v3 单腿做多'],[
 ['信号','近月减远月的收盘价差连涨3次','相同的价差连涨信号'],
 ['入场','次日开盘多近月1手 空远月1手','分别测试只买近月或只买远月1手'],
 ['止盈止损对象','组合价差相对入场价差','所买合约相对其入场价格'],
 ['默认阈值','止盈5点 止损3点','止盈5点 止损3点'],
 ['退出时点','日收盘触发 次日开盘退出','日收盘触发 次日开盘退出'],
 ['假设往返成本','两腿共400元','单腿200元']],[1.25,2.85,2.9])
p('连续上涨3次至少需要4个日收盘价。价差相等或下降会重置计数；持仓期间不累计新信号，不加仓；退出当天不重新累计。v3两个策略分别运行，退出时间不同可能导致后续入场日期和笔数不同，因此不是固定同一持有期的逐笔比较。')
sub('组合收益必须合并计算')
p('v2的近月多单赚取近月平仓价减开仓价；远月空单赚取远月开仓价减平仓价。两项相加，再乘1000并扣除成本，才是每组净收益。空单亏损并不代表组合亏损，只要近月多单赚得更多，组合仍可盈利。')
p('v3只保留所买合约的多单收益，不再用另一腿抵消原油价格方向风险。两个方案的回测报告不应相加成为一个策略的收益。')
sub('流动性过滤的含义')
p('原始规则不增加流动性限制。研究对照版仅在信号当天检查：远月持仓量至少1000手，近月和远月各自当日成交量至少100手。筛选不使用下一交易日的信息，也不阻止已有持仓退出。这一附加条件尚未写入MC v3信号。')
sub('盯市与异常数据处理')
p('期末未平仓按最后收盘价估值，只扣已发生的半程入场成本，不假装已经平仓。最大回撤按各年份内部日收盘净值计算，包含持仓浮动盈亏；它不是盘中最大回撤，也没有包含保证金、资金占用或强平。')
p('Python研究在缺少开盘价时取消入场或延后退出，缺少持仓收盘价时沿用上一有效估值。MC对数据或持仓异常会报错。原始规则中存在缺价处理，不能宣称这类区间与MC完全一致；默认参数过滤版未出现缺开盘价的入场或退出异常。')

page();h('v2 双腿价差交易结果')
p('原始规则共完成24笔交易，已平仓净赚2700元；加上2024年样本末未平仓净盯市亏损2700元，合计约为零。过滤版完成14笔，合计净赚10600元，净胜率42.9%，七个年份中四盈三亏。')
table(['合约年份','原始净收益','过滤净收益','过滤日线回撤'],[
 [y,money(v2['results']['raw'][y]['total_mark']),money(r['total_mark']),money(r['mdd'])]
 for y,r in v2['results']['liquid'].items()],[1.25,1.85,1.85,2.05])
p('上表金额单位为元。2026年原始规则7笔净亏3000元，过滤版4笔净赚4500元，但过滤版该年日线最大回撤为20900元。')
sub('一笔盈利交易及收益集中度')
table(['2026年8月18日至9月1日','收益金额'],[
 ['近月多单毛收益','70,200元'],['远月空单毛收益','−55,100元'],
 ['组合毛收益','15,100元'],['扣400元成本后的组合净收益','14,700元']],[4.4,2.6])
p('这一笔盈利大于过滤版七个年份合计利润。仅作集中度检查，若去掉这一笔，2026年结果为亏10200元，全部年份合计亏4100元。不能据此删除交易，但它说明收益高度依赖少数行情。')
sub('参数与成本敏感性')
p('过滤版将连涨次数从3改为2，合计净盯市变为亏2200元；改为5次则亏6700元。止盈从5点改为4点，合计亏11600元。默认参数下，组合往返成本从400元提高到800元，合计净收益由10600元降至5000元。参数不稳定与样本偏小仍是主要问题。')

page();h('v3 近月做多与远月做多结果')
p('按原始规则，近月49笔合计净赚19600元，远月41笔合计净亏34900元。增加相同流动性过滤后，两者各22笔：近月净赚65300元、净胜率63.6%；远月净亏20700元、净胜率54.5%。胜率较高也不一定盈利，单笔盈亏幅度同样重要。')
table(['合约年份','近月过滤净收益','远月过滤净收益','近月过滤回撤'],[
 [y,money(r['total_mark']),money(v3['results']['liquid']['far'][y]['total_mark']),money(r['mdd'])]
 for y,r in v3['results']['liquid']['near'].items()],[1.25,1.9,1.9,1.95])
p('上表金额单位为元。近月过滤版六个年份盈利，但2026年日线最大回撤94800元，明显高于该合约样本最终净收益26700元。取消空单后，利润增加并不等于风险调整后的表现更好。')
sub('2026年合约样本的筛选差异')
table(['测试口径','近月做多净收益','远月做多净收益'],[
 ['原始规则','−51,600元','−94,000元'],['流动性过滤','26,700元','−21,600元']],[2.4,2.3,2.3])
p('近月是否盈利在加入过滤后发生改变。过滤规则只降低部分早期不活跃行情的影响，不能证明筛选后的利润一定可以实际成交。当前MC v3代码对应原始规则，不是过滤版。')
sub('近远月优势并非所有参数都成立')
p('在过滤版中，将连涨次数改为2，近月合计净赚16700元，远月净赚75100元，优劣发生反转。保持连涨3次、仅将止盈改为4点或6点，近月仍分别净赚63900元、53400元；这只能说明部分邻近参数有正收益，不能代替独立样本验证。')
p('因此，本次结果支持继续研究近月单腿做多，但不支持直接认定“近月总比远月好”，也没有证明价差连涨能独立预测单腿未来涨跌。尚需与无信号基准及固定共同持有期对照。')

page();h('风险边界与后续验证')
sub('止损阈值不是最大损失')
p('本研究的止损3点，是在日收盘时检查亏损是否达到阈值，再于下一交易日开盘退出。隔夜跳空、信号时点已经超出阈值、单腿价格波动都可能使实际亏损远超3000元。v2按价差止损，v3按单腿止损，二者不能沿用相同风险预期。')
sub('日线不能证明同时可成交')
p('两份期货的日开盘或收盘价可能来自不同时间，远月不活跃时尤其明显。v2以两腿日开盘价假设同时成交，可能高估可交易利润；v3虽只有一腿下单，其价差信号仍受不同步收盘报价影响。回测没有模拟盘口容量、涨跌停排队、拒单、单腿成交、保证金不足或到期强平。')
sub('下一阶段应验证什么')
p('优先取得同一时间戳的两腿分钟行情或买卖盘口，用可成交买卖价重算信号与成交。保留当前默认参数作为待检验规则，另设未参与选择的时间区间；同时比较无信号基准、相同入场日和固定持有期，区分原油方向行情与价差信号的贡献。')
p('在考虑实盘前，还需加入最大持有期、交割前退出、实际成本和保证金约束，并检验异常行情下的退出能力。这些都是后续计划，不是本次已完成的测试。当前成果属于历史研究，未启用自动交易。')
sub('MultiCharts使用边界')
p('v3使用同一份Signal代码：近月方案将SC2610设为Data1、SC2611为Data2，NearIsData1=True；远月方案对调数据，参数为False。两图均为1日线，代码只做多Data1。不得叠加原v2信号误当成同一策略。[3]')
p('Python共25项测试通过，其中v3新增8项，覆盖时序、单腿方向、止损对象、成本及部分数据异常。尚未在MC实际编译或核对交易报告；测试通过不能被理解为实盘可用。')
sub('资料来源与复现文件')
for txt in [
 '[1] 上期能源原油期货标准合约，每手1000桶。https://www.ine.cn/products/futures/energyandchemical/sc_f/standard_sc_f/202312/t20231205_802540.html',
 '[2] 上期能源公开日行情接口，示例日期2026年9月8日。https://www.ine.cn/data/tradedata/future/dailydata/kx20260908.dat',
 '[3] MultiCharts订单语义说明。https://www.multicharts.cn/post?idno=169',
 '复现目录为TASK_0004_SC跨月价差研究。数据文件：sc_oct_nov_panel_2020_2026.csv；研究脚本：backtest_v2_pair.py与backtest_v3_long.py；逐笔及每日净值见output目录对应JSON。']:
 para=p(txt)
 for run in para.runs:run.font.size=Pt(9)

footer=s.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=footer.add_run();field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');r._r.addnext(field)
out=ROOT/'output/原油跨月价差与单腿做多策略研究.docx'
d.save(out)
print(out)
